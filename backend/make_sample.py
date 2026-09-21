"""Generate a synthetic PAIMANA-style monthly monitoring report (Excel) for demos.

Layout mimics ministry exports: title rows above the header, ₹ crore columns,
mixed date formats and a Total row at the bottom — so the parser gets exercised.
"""

import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUT = Path(__file__).parent / "sample_data" / "sample_paimana_report.xlsx"
AS_OF = date.today()
rng = random.Random(26103)

SECTORS = {
    # sector: (ministry, agencies, typical slowness, project templates)
    "Railways": ("Ministry of Railways", ["RVNL", "IRCON", "Northern Railway", "East Coast Railway", "DFCCIL"], 1.45,
                 ["New Line {a}–{b}", "Doubling {a}–{b}", "Gauge Conversion {a}–{b}", "Electrification {a}–{b}"]),
    "Road Transport & Highways": ("Ministry of Road Transport & Highways", ["NHAI", "NHIDCL"], 1.3,
                                  ["4-laning of NH-{n} {a}–{b}", "{a}–{b} Expressway Pkg-{n}", "Bypass at {a}"]),
    "Power": ("Ministry of Power", ["NTPC", "NHPC", "PGCIL", "SJVN", "THDC"], 1.35,
              ["{a} Super Thermal Power Project Stage-{r}", "{a} Hydro Electric Project", "{a}–{b} 765kV Transmission Line"]),
    "Petroleum & Natural Gas": ("Ministry of Petroleum & Natural Gas", ["IOCL", "BPCL", "GAIL", "ONGC", "HPCL"], 1.2,
                                ["{a}–{b} Natural Gas Pipeline", "{a} Refinery Expansion", "{a} LPG Import Terminal"]),
    "Coal": ("Ministry of Coal", ["Coal India Ltd", "NLC India", "SCCL"], 1.25,
             ["{a} Opencast Expansion", "{a} First Mile Connectivity", "{a} Coal Washery"]),
    "Urban Development": ("Ministry of Housing & Urban Affairs", ["DMRC", "MMRCL", "BMRCL", "NCRTC"], 1.35,
                          ["{a} Metro Phase-{r}", "{a}–{b} RRTS Corridor"]),
    "Water Resources": ("Ministry of Jal Shakti", ["NWDA", "WAPCOS", "State Irrigation Dept"], 1.6,
                        ["{a} Irrigation Project", "{a} Dam Rehabilitation", "{a}–{b} River Link Canal"]),
    "Civil Aviation": ("Ministry of Civil Aviation", ["Airports Authority of India"], 1.15,
                       ["New Terminal Building at {a} Airport", "{a} Greenfield Airport"]),
    "Shipping & Ports": ("Ministry of Ports, Shipping & Waterways", ["Paradip Port", "Deendayal Port", "JNPA", "IWAI"], 1.25,
                         ["Deep Draught Berth at {a}", "{a} Inland Waterway Terminal"]),
    "Telecommunications": ("Department of Telecommunications", ["BSNL", "BBNL"], 1.4,
                           ["BharatNet Phase-{r} {a} Circle", "4G Saturation {a} Circle"]),
    "Steel": ("Ministry of Steel", ["SAIL", "NMDC", "RINL"], 1.2,
              ["{a} Steel Plant Modernisation", "{a} Pellet Plant"]),
    "Health & Family Welfare": ("Ministry of Health & Family Welfare", ["HSCC", "CPWD"], 1.2,
                                ["AIIMS {a}", "{a} Critical Care Block"]),
}
PLACES = [("Patna", "Bihar", "Gaya"), ("Ranchi", "Jharkhand", "Dhanbad"), ("Bhubaneswar", "Odisha", "Cuttack"),
          ("Guwahati", "Assam", "Silchar"), ("Lucknow", "Uttar Pradesh", "Kanpur"), ("Nagpur", "Maharashtra", "Wardha"),
          ("Surat", "Gujarat", "Vadodara"), ("Jaipur", "Rajasthan", "Ajmer"), ("Bhopal", "Madhya Pradesh", "Itarsi"),
          ("Raipur", "Chhattisgarh", "Bilaspur"), ("Dehradun", "Uttarakhand", "Rishikesh"),
          ("Shimla", "Himachal Pradesh", "Kalka"), ("Chennai", "Tamil Nadu", "Vellore"), ("Kochi", "Kerala", "Thrissur"),
          ("Vizag", "Andhra Pradesh", "Vizianagaram"), ("Hyderabad", "Telangana", "Warangal"),
          ("Mangaluru", "Karnataka", "Udupi"), ("Siliguri", "West Bengal", "Jalpaiguri"),
          ("Jammu", "Jammu & Kashmir", "Udhampur"), ("Imphal", "Manipur", "Jiribam"), ("Ludhiana", "Punjab", "Jalandhar"),
          ("Rohtak", "Haryana", "Hisar")]
REASONS = ["Land acquisition", "Forest clearance", "Environment clearance", "Contractual issues",
           "Funding constraints", "Utility shifting", "Law & order", "Rehabilitation & resettlement",
           "Design changes", "Monsoon / geological surprises"]


def month_str(d: date) -> str:
    return d.strftime("%b-%Y")


def make_project(i: int, sector: str | None = None, start_days=(200, 3400)) -> list:
    sector = sector or rng.choices(list(SECTORS), weights=[16, 18, 12, 9, 7, 8, 7, 4, 5, 4, 4, 6])[0]
    ministry, agencies, slow, templates = SECTORS[sector]
    a, state, b = rng.choice(PLACES)
    name = rng.choice(templates).format(a=a, b=b, n=rng.randint(2, 99), r=rng.choice(["I", "II", "III"]))

    original = round(rng.choice([rng.uniform(150, 1000), rng.uniform(1000, 5000), rng.uniform(5000, 25000)]), 2)
    start = AS_OF - timedelta(days=rng.randint(*start_days))
    planned = rng.randint(24, 72) * 30
    orig_end = start + timedelta(days=planned)

    # True execution pace (hidden from the agency's own report).
    pace = max(0.85, rng.lognormvariate(0, 0.28) * slow)
    elapsed = (AS_OF - start).days
    progress = min(100.0, max(1.0, elapsed / (planned * pace) * 100 + rng.gauss(0, 4)))
    progress = round(progress, 1)

    # Agencies revise dates late and partially.
    if progress >= 100:
        antic = start + timedelta(days=round(planned * pace))
    else:
        honesty = rng.uniform(0.1, 1.0) if elapsed < planned else rng.uniform(0.5, 1.0)
        antic = orig_end + timedelta(days=round(planned * (pace - 1) * honesty))
        antic = max(antic, AS_OF + timedelta(days=60)) if antic < AS_OF else antic
        if rng.random() < 0.25 and elapsed < planned * 0.8:
            antic = orig_end  # "on schedule" as reported

    overrun = max(0.0, (pace - 1) * rng.uniform(0.2, 0.6) + rng.gauss(0.02, 0.05))
    revised = round(original * (1 + (overrun if rng.random() < 0.7 else 0)), 2)
    burn = rng.uniform(0.85, 1.1) if rng.random() < 0.8 else rng.uniform(1.25, 1.6)
    spent = round(min(revised * 1.05, revised * progress / 100 * burn), 2)

    reasons = ", ".join(rng.sample(REASONS, rng.randint(1, 2))) if pace > 1.3 and rng.random() < 0.85 else ""
    status = "Completed" if progress >= 100 else "Delayed" if antic > orig_end else "On schedule"
    pid = f"{ministry.split()[-1][:3].upper()}-{2026}-{i:04d}"
    return [i, pid, name, ministry, sector, state, rng.choice(agencies), original, revised, spent,
            start, orig_end, antic, progress, status, reasons]


def _style_header(ws, row: int, color: str = "1F3A5F"):
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=color)


def paimana_report(path: Path):
    """All-sector ministry export: title rows above the header, Total row, month-year anticipated dates."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Project Report"
    ws.append(["Project Monitoring Report — Central Sector Infrastructure Projects (₹150 crore & above)"])
    ws.append([f"Synthetic demo data · Report month: {AS_OF.strftime('%B %Y')} · Costs in ₹ crore"])
    ws.append([])
    ws.append(["S.No.", "Project ID", "Project Name", "Ministry/Department", "Sector", "State",
               "Implementing Agency", "Original Cost (₹ Cr)", "Latest Approved / Revised Cost (₹ Cr)",
               "Cumulative Expenditure (₹ Cr)", "Date of Approval", "Original Date of Commissioning",
               "Anticipated Date of Commissioning", "Physical Progress (%)", "Status", "Reasons for Delay"])
    _style_header(ws, 4)
    ws["A1"].font = Font(bold=True, size=13)
    rows = [make_project(i + 1) for i in range(180)]
    for r in rows:
        ws.append(r[:12] + [month_str(r[12])] + r[13:])
    ws.append(["", "", "Total", "", "", "", "", round(sum(r[7] for r in rows), 2),
               round(sum(r[8] for r in rows), 2), round(sum(r[9] for r in rows), 2)])
    for row in ws.iter_rows(min_row=5, min_col=11, max_col=12):
        for c in row:
            c.number_format = "DD-MM-YYYY"
    for col, width in zip("ABCDEFGHIJKLMNOP", [6, 14, 44, 34, 24, 16, 20, 14, 18, 16, 14, 16, 16, 12, 12, 36]):
        ws.column_dimensions[col].width = width
    wb.save(path)
    return len(rows)


RAIL_KINDS = ["New Line", "Doubling", "Gauge Conversion", "Electrification"]


def railways_review(path: Path):
    """Single-ministry workbook: a cover sheet first, different header wording, progress as a 0–1 fraction,
    dd/mm/yyyy text dates, and 'Bottlenecks' instead of reasons."""
    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    for line in ["Ministry of Railways — Zonal Project Review", f"Review date: {AS_OF.strftime('%d/%m/%Y')}",
                 "Synthetic demo data", "", "Sheet 'Projects' lists all sanctioned works above ₹150 crore."]:
        cover.append([line])
    ws = wb.create_sheet("Projects")
    ws.append(["Proj. Code", "Name of the Project", "Category", "Zone / State", "Executing Agency",
               "Sanctioned Cost (Rs. Cr.)", "Anticipated Cost (Rs. Cr.)", "Exp. Incurred till date (Rs. Cr.)",
               "Sanction Date", "Target Date", "Revised DOC", "% Complete", "Bottlenecks"])
    _style_header(ws, 1, "7A1F1F")
    n = 70
    for i in range(n):
        r = make_project(i + 1, sector="Railways")
        kind = next((k for k in RAIL_KINDS if r[2].startswith(k)), "New Line")
        ws.append([f"RB/{r[10].year}/{i + 1:03d}", r[2], f"Railways - {kind}", r[5], r[6], r[7], r[8], r[9],
                   r[10].strftime("%d/%m/%Y"), r[11].strftime("%d/%m/%Y"), r[12].strftime("%d/%m/%Y"),
                   round(r[13] / 100, 3), r[15].replace(", ", "; ")])
    for col, width in zip("ABCDEFGHIJKLM", [14, 40, 26, 16, 18, 14, 14, 16, 12, 12, 12, 10, 40]):
        ws.column_dimensions[col].width = width
    wb.save(path)
    return n


def state_export_csv(path: Path):
    """Flat CSV from a state dashboard: generic headers, ISO dates, progress as '45%' text."""
    import csv
    n = 90
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["ID", "Title", "Department", "Sector", "Location", "Approved Cost", "Current Cost", "Spent",
                    "Start Date", "Planned Completion", "Expected Completion", "Progress", "Current Status",
                    "Remarks"])
        for i in range(n):
            r = make_project(i + 1)
            w.writerow([f"ST-{i + 1:04d}", r[2], r[3], r[4], r[5], f"{r[7]:,.2f}", f"{r[8]:,.2f}", f"{r[9]:,.2f}",
                        r[10].isoformat(), r[11].isoformat(), r[12].isoformat(), f"{r[13]:.0f}%", r[14], r[15]])
    return n


def early_stage_minimal(path: Path):
    """Minimal sheet for newly sanctioned projects: no expenditure, no anticipated dates, no reasons —
    shows how the dashboard degrades gracefully when CUF fields are missing."""
    wb = Workbook()
    ws = wb.active
    ws.title = "New Sanctions"
    ws.append(["Project", "Sector", "State", "Cost (₹ Cr)", "Start Date", "Completion Date", "Progress %"])
    _style_header(ws, 1, "1F5F3A")
    n = 45
    for i in range(n):
        r = make_project(i + 1, start_days=(90, 900))
        ws.append([r[2], r[4], r[5], r[7], r[10], r[11], r[13]])
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=6):
        for c in row:
            c.number_format = "DD-MMM-YYYY"
    for col, width in zip("ABCDEFG", [44, 24, 16, 12, 14, 16, 10]):
        ws.column_dimensions[col].width = width
    wb.save(path)
    return n


SAMPLES = [
    ("sample_paimana_report.xlsx", paimana_report),
    ("railways_zone_review.xlsx", railways_review),
    ("state_projects_export.csv", state_export_csv),
    ("early_stage_sanctions.xlsx", early_stage_minimal),
]


def main():
    OUT.parent.mkdir(exist_ok=True)
    for name, build in SAMPLES:
        count = build(OUT.parent / name)
        print(f"Wrote {count} projects -> {OUT.parent / name}")


if __name__ == "__main__":
    main()
